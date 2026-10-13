#!/usr/bin/env python3
"""Check requested IP ownership in an Excel workbook's Envir worksheet.

A clean worksheet alone does NOT prove an address is available on the live network.
"""
from __future__ import annotations
import argparse
import ipaddress
import sys
from pathlib import Path
from typing import Any

try:
    from openpyxl import Workbook, load_workbook
except ImportError:
    print("ERROR: install Python dependencies using: python -m pip install -r requirements.txt", file=sys.stderr)
    raise SystemExit(2)

HEADERS = [
    "MachineName", "Role", "Hypervisor", "IP_Address", "Subnet_Prefix",
    "Gateway", "DNS", "VLAN_or_Network", "VM_ID", "OS", "Status",
    "Owner", "Last_Checked", "Notes",
]
NAME_KEYS = {"machinename", "hostname", "computername", "vmname", "name", "servername", "server"}
IP_KEYS = {"ipaddress", "ip", "ipv4", "ipaddressv4", "ipv4address"}

def norm(value: Any) -> str:
    return "".join(char.lower() for char in str(value or "").strip() if char.isalnum())

def normal_ip(value: Any) -> str:
    raw = str(value or "").strip().split("/")[0]
    try:
        return str(ipaddress.IPv4Address(raw)) if raw else ""
    except ipaddress.AddressValueError:
        return ""

def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", required=True)
    parser.add_argument("--target", action="append", default=[], metavar="NAME=IP")
    parser.add_argument("--create-template", action="store_true")
    args = parser.parse_args()
    path = Path(args.workbook).expanduser()

    if args.create_template:
        if path.exists():
            print(f"ERROR: refusing to overwrite {path}", file=sys.stderr)
            return 2
        path.parent.mkdir(parents=True, exist_ok=True)
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Envir"
        sheet.append(HEADERS)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = f"A1:N1"
        widths = [22,18,18,18,16,16,18,22,12,16,14,18,18,34]
        for index, width in enumerate(widths, start=1):
            sheet.column_dimensions[chr(64 + index)].width = width
        workbook.save(path)
        print(f"Created blank template at {path}. It does not confirm IP availability; fill it before checks.")
        return 0

    if not path.is_file():
        print(f"ERROR: workbook not found: {path}", file=sys.stderr)
        return 2
    if not args.target:
        print("ERROR: specify at least one --target NAME=IP", file=sys.stderr)
        return 2

    targets = {}
    for item in args.target:
        if "=" not in item:
            print(f"ERROR: invalid target {item!r}; expected NAME=IP", file=sys.stderr)
            return 2
        name, raw_ip = item.split("=", 1)
        ip = normal_ip(raw_ip)
        if not name.strip() or not ip:
            print(f"ERROR: invalid target {item!r}", file=sys.stderr)
            return 2
        targets[name.strip().casefold()] = ip

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if "Envir" not in workbook.sheetnames:
            print("ERROR: workbook has no worksheet named Envir.", file=sys.stderr)
            return 2
        rows = workbook["Envir"].iter_rows(values_only=True)
        header = next(rows, None)
        if not header:
            print("ERROR: Envir has no header row.", file=sys.stderr)
            return 2
        headers = [norm(value) for value in header]
        ip_index = next((i for i, value in enumerate(headers) if value in IP_KEYS), None)
        name_index = next((i for i, value in enumerate(headers) if value in NAME_KEYS), None)
        if ip_index is None:
            print("ERROR: add an IP column, e.g. IP_Address or IPv4.", file=sys.stderr)
            return 2

        records = 0
        conflicts = []
        for row_number, row in enumerate(rows, start=2):
            if not row or not any(value is not None and str(value).strip() for value in row):
                continue
            records += 1
            current_ip = normal_ip(row[ip_index] if ip_index < len(row) else "")
            if not current_ip:
                continue
            current_name = str(row[name_index] or "").strip().casefold() if name_index is not None and name_index < len(row) else ""
            for target_name, target_ip in targets.items():
                if current_ip == target_ip and current_name != target_name:
                    conflicts.append(f"row {row_number}: {current_ip} belongs to {current_name or '(blank name)'}, not {target_name}")
    finally:
        workbook.close()

    if records == 0:
        print("BLOCKED: Envir has no inventory records; cannot check ownership.")
        return 3
    if conflicts:
        print("FAIL: Envir ownership conflicts found:")
        for message in conflicts:
            print(f"- {message}")
        return 1

    print(f"PASS: no conflicting Envir ownership found for {len(targets)} address(es).")
    for name, ip in targets.items():
        print(f"- {name}: {ip} (worksheet check only)")
    print("REQUIRED: verify DHCP/router/IPAM reservations and live network conflicts before assigning IPs.")
    return 0

if __name__ == "__main__":
    sys.exit(main())
